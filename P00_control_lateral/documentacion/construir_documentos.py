"""
Construye los tres documentos del registro del software con pandoc.

Código auxiliar nuevo, no modifica ningún script existente.

Entradas
    fuentes/descripcion.md, fuentes/manual_tecnico.md, fuentes/manual_usuario.md
    datos_portada.json con los datos que entrega el investigador
    figuras/ y capturas/

Marcas que se resuelven antes de pandoc
    {{CAMPO}}          dato de portada
    [[fig:id]]         número de figura por orden de primera aparición
    [[tab:id]]         número de tabla por orden de primera aparición
    [[eq:id]]          número de ecuación por orden de primera aparición
    [[ref:a,b]]        cita numerada por orden de primera aparición, entre corchetes
    [[REFERENCIAS]]    lista de referencias en el orden de las citas
    [[TOC]]            tabla de contenido de Word
    [[SALTO]]          salto de página

Una imagen de capturas/ que todavía no existe se sustituye por un aviso de
captura pendiente, sin inventar la imagen.

Salida en documentacion/, un .docx por documento, y un informe de auditoría
de marcas y de reglas de redacción en documentacion/auditoria_documentos.txt.

Uso desde la raíz del repositorio.
    python P00_control_lateral/documentacion/construir_documentos.py
Después, para llenar la tabla de contenido con Word y dejar un PDF de revisión,
    powershell -File P00_control_lateral/documentacion/actualizar_con_word.ps1 P00_control_lateral/documentacion/revision_pdf
Llamar a Word desde este script lo deja colgado, verificado el 2026-09-23, por eso es un paso aparte.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DIR = Path(__file__).resolve().parent
FUENTES = DIR / "fuentes"

DOCUMENTOS = [
    ("descripcion.md", "Descripcion_del_Software_{nombre}.docx", "Descripción del software"),
    ("manual_tecnico.md", "Manual_Tecnico_{nombre}.docx", "Manual técnico"),
    ("manual_usuario.md", "Manual_de_Usuario_{nombre}.docx", "Manual de usuario"),
]

# Referencias con DOI resuelto el 2026-09-23 en Crossref y, para arXiv, en DataCite.
REFERENCIAS = {
    "zhang2025": "W. Zhang, J. Wang y T. Pang, Research on the Performance of Vehicle Lateral Control Algorithm "
                 "Based on Vehicle Speed Variation, World Electric Vehicle Journal, vol. 16, no. 5, art. 259, 2025. "
                 "doi 10.3390/wevj16050259",
    "kong2024": "X. Kong, C. Fisher y B. Evans, A critical evaluation of Pure Pursuit, MPC and MPCC, balancing "
                "simplicity, performance and constraints, MATEC Web of Conferences, vol. 406, art. 04013, 2024. "
                "doi 10.1051/matecconf/202440604013",
    "lee2023": "J. Lee y S. Yim, Comparative Study of Path Tracking Controllers on Low Friction Roads for Autonomous "
               "Vehicles, Machines, vol. 11, no. 3, art. 403, 2023. doi 10.3390/machines11030403",
    "bousskoul2025": "A. Bousskoul, I. Ouachtouk y A. Ait Elmahjoub, Control strategies for autonomous vehicle path "
                     "tracking, a comparative study of PID, Pure Pursuit, and Stanley methods, EPJ Web of Conferences, "
                     "vol. 330, art. 06001, 2025. doi 10.1051/epjconf/202533006001",
    "remonda2024": "A. Remonda, N. Hansen, A. Raji, N. Musiu, M. Bertogna, E. Veas y X. Wang, A Simulation Benchmark "
                   "for Autonomous Racing with Large Scale Human Data, arXiv, 2024. doi 10.48550/arXiv.2407.16680",
    "bockman2024": "J. Bockman, M. Howe, A. Orenstein y F. Dayoub, AARK, An Open Toolkit for Autonomous Racing "
                   "Research, arXiv, 2024. doi 10.48550/arXiv.2410.00358",
    "ramlall2025": "P. Ramlall, E. Jones y S. Roy, Development of a Networked Multi Participant Driving Simulator with "
                   "Synchronized EEG and Telemetry for Traffic Research, Systems, vol. 13, no. 7, art. 564, 2025. "
                   "doi 10.3390/systems13070564",
    "hoffmann2007": "G. Hoffmann, C. Tomlin, M. Montemerlo y S. Thrun, Autonomous Automobile Trajectory Tracking for "
                    "Off Road Driving, Controller Design, Experimental Validation and Racing, en 2007 American Control "
                    "Conference, 2007, pp. 2296 a 2301. doi 10.1109/ACC.2007.4282788",
    "kong2015": "J. Kong, M. Pfeiffer, G. Schildbach y F. Borrelli, Kinematic and dynamic vehicle models for "
                "autonomous driving control design, en 2015 IEEE Intelligent Vehicles Symposium, 2015, pp. 1094 a "
                "1099. doi 10.1109/IVS.2015.7225830",
    "stellato2020": "B. Stellato, G. Banjac, P. Goulart, A. Bemporad y S. Boyd, OSQP, an operator splitting solver for "
                    "quadratic programs, Mathematical Programming Computation, vol. 12, no. 4, pp. 637 a 672, 2020. "
                    "doi 10.1007/s12532-020-00179-2",
    "skogestad2003": "S. Skogestad, Simple analytic rules for model reduction and PID controller tuning, Journal of "
                     "Process Control, vol. 13, no. 4, pp. 291 a 309, 2003. doi 10.1016/S0959-1524(02)00062-8",
}

AZUL = RGBColor(0x00, 0x3E, 0x7E)
PAGINA = "<w:p><w:r><w:br w:type=\"page\"/></w:r></w:p>"
TOC = ("<w:p><w:pPr><w:pStyle w:val=\"TOCHeading\"/></w:pPr><w:r><w:t>Tabla de contenido</w:t></w:r></w:p>"
       "<w:p><w:r><w:fldChar w:fldCharType=\"begin\" w:dirty=\"true\"/></w:r>"
       "<w:r><w:instrText xml:space=\"preserve\"> TOC \\o \"1-3\" \\h \\z \\u </w:instrText></w:r>"
       "<w:r><w:fldChar w:fldCharType=\"separate\"/></w:r>"
       "<w:r><w:t>La tabla de contenido se actualiza al abrir el documento.</w:t></w:r>"
       "<w:r><w:fldChar w:fldCharType=\"end\"/></w:r></w:p>")


def cargar_datos():
    with open(DIR / "datos_portada.json", encoding="utf-8") as f:
        datos = json.load(f)
    return {k: v for k, v in datos.items() if not k.startswith("_")}


def resolver(texto, datos, avisos):
    for k, v in datos.items():
        texto = texto.replace("{{" + k + "}}", v)
    for faltante in sorted(set(re.findall(r"\{\{(\w+)\}\}", texto))):
        avisos.append(f"campo de portada sin definir {faltante}")

    numeros = {"fig": {}, "tab": {}, "eq": {}, "ref": {}}
    for m in re.finditer(r"\[\[(fig|tab|eq|ref):([\w,\s]+)\]\]", texto):
        tipo = m.group(1)
        for clave in [c.strip() for c in m.group(2).split(",")]:
            if clave not in numeros[tipo]:
                numeros[tipo][clave] = len(numeros[tipo]) + 1

    def sustituir(m):
        tipo = m.group(1)
        claves = [c.strip() for c in m.group(2).split(",")]
        if tipo == "ref":
            for c in claves:
                if c not in REFERENCIAS:
                    avisos.append(f"cita sin referencia {c}")
            return "[" + ", ".join(str(numeros["ref"][c]) for c in claves) + "]"
        return ", ".join(str(numeros[tipo][c]) for c in claves)

    texto = re.sub(r"\[\[(fig|tab|eq|ref):([\w,\s]+)\]\]", sustituir, texto)

    lista = sorted(numeros["ref"].items(), key=lambda kv: kv[1])
    bloque = "\n\n".join(f"[{n}] {REFERENCIAS.get(c, 'REFERENCIA NO ENCONTRADA ' + c)}" for c, n in lista)
    texto = texto.replace("[[REFERENCIAS]]", bloque)
    texto = texto.replace("[[TOC]]", "```{=openxml}\n" + TOC + "\n```")
    texto = texto.replace("[[SALTO]]", "```{=openxml}\n" + PAGINA + "\n```")

    def imagen(m):
        pie, ruta, atributos = m.group(1), m.group(2), m.group(3) or ""
        if (DIR / ruta).exists():
            return m.group(0)
        avisos.append(f"imagen pendiente {ruta}")
        return (f"::: {{custom-style=\"Pendiente\"}}\n[CAPTURA PENDIENTE DE EJECUCIÓN, {Path(ruta).name}]\n:::\n\n"
                f"::: {{custom-style=\"Image Caption\"}}\n{pie}\n:::")

    texto = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)(\{[^}]*\})?", imagen, texto)
    return texto, numeros


def referencia_docx(ruta):
    """Plantilla de estilos propia, derivada de la de pandoc."""
    datos = subprocess.run(["pandoc", "-o", str(ruta), "--print-default-data-file", "reference.docx"],
                           capture_output=True)
    if datos.returncode != 0:
        raise RuntimeError(datos.stderr.decode(errors="ignore"))
    doc = Document(ruta)
    st = doc.styles

    def fuente(estilo, tam, color=None, negrita=None, cursiva=None):
        f = st[estilo].font
        f.name = "Calibri"
        f.size = Pt(tam)
        if color is not None:
            f.color.rgb = color
        if negrita is not None:
            f.bold = negrita
        if cursiva is not None:
            f.italic = cursiva
        rpr = st[estilo].element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), "Calibri")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]

    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if nombre in [s.name for s in st]:
            fuente(nombre, 11, RGBColor(0, 0, 0))
    for nombre in ("Body Text", "First Paragraph"):
        pf = st[nombre].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
    fuente("Heading 1", 16, AZUL, True, False)
    fuente("Heading 2", 13, AZUL, True, False)
    fuente("Heading 3", 12, AZUL, True, False)
    for nombre, antes, despues in (("Heading 1", 18, 8), ("Heading 2", 12, 6), ("Heading 3", 10, 4)):
        st[nombre].paragraph_format.space_before = Pt(antes)
        st[nombre].paragraph_format.space_after = Pt(despues)
    fuente("Title", 28, AZUL, True, False)
    st["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente("Subtitle", 14, RGBColor(0x33, 0x33, 0x33), False, False)
    st["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for nombre in ("Image Caption", "Table Caption"):
        fuente(nombre, 10, RGBColor(0x22, 0x22, 0x22), False, True)
        st[nombre].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st[nombre].paragraph_format.space_after = Pt(10)
    st["Captioned Figure"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Figure"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Pendiente" not in [s.name for s in st]:
        p = st.add_style("Pendiente", 1)
        p.base_style = st["Body Text"]
        p.font.color.rgb = RGBColor(0xA0, 0x10, 0x10)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Cubierta" not in [s.name for s in st]:
        p = st.add_style("Cubierta", 1)
        p.base_style = st["Body Text"]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)
    doc.save(ruta)


def campo_pagina(parrafo):
    for tipo, texto in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        run = parrafo.add_run()
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        elif texto == " PAGE ":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = texto
            run._r.append(it)
        else:
            run.text = texto


def bordes_tabla(tabla, visibles=True):
    tblpr = tabla._tbl.tblPr
    previo = tblpr.find(qn("w:tblBorders"))
    if previo is not None:
        tblpr.remove(previo)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single" if visibles else "nil")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "8C8C8C")
        bordes.append(e)
    tblpr.append(bordes)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def posproceso(ruta, nombre, titulo_doc):
    doc = Document(ruta)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    cab = sec.header.paragraphs[0]
    cab.text = f"{nombre}, {titulo_doc}"
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in cab.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    campo_pagina(pie)
    for i, tabla in enumerate(doc.tables):
        portada = i == 0
        bordes_tabla(tabla, visibles=not portada)
        for fila_i, fila in enumerate(tabla.rows):
            for celda in fila.cells:
                for p in celda.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(9.5 if not portada else 11)
                if fila_i == 0 and not portada:
                    sombrear(celda, "DCE8F5")
    # Word llena la tabla de contenido al abrir el documento, si no se corrió actualizar_con_word.ps1.
    act = OxmlElement("w:updateFields")
    act.set(qn("w:val"), "true")
    doc.settings.element.append(act)
    doc.save(ruta)


def auditar(nombre_md, texto_final):
    """Reglas de redacción de AGENTS.md sobre el texto que ve el lector."""
    hallazgos = []
    cuerpo = re.sub(r"```.*?```", "", texto_final, flags=re.S)
    cuerpo = re.sub(r"`[^`]*`", "", cuerpo)
    cuerpo = re.sub(r"\$\$.*?\$\$", "", cuerpo, flags=re.S)
    cuerpo = re.sub(r"\$[^$]*\$", "", cuerpo)
    cuerpo = re.sub(r"!\[[^\]]*\]\([^)]*\)(\{[^}]*\})?", "", cuerpo)
    cuerpo = re.sub(r"\{[^}]*\}", "", cuerpo)
    for n, linea in enumerate(cuerpo.splitlines(), 1):
        l = linea.strip()
        if not l or l.startswith("|--") or re.fullmatch(r"[|:\- ]+", l):
            continue
        if l.startswith(": "):
            l = l[2:]
        if "—" in l or "–" in l:
            hallazgos.append(f"{nombre_md} línea {n} guion largo")
        if ";" in l:
            hallazgos.append(f"{nombre_md} línea {n} punto y coma")
        if re.search(r"[\"“”«»]", l):
            hallazgos.append(f"{nombre_md} línea {n} comillas")
        if "**" in l or "__" in l:
            hallazgos.append(f"{nombre_md} línea {n} negrilla")
        sin_url = re.sub(r"https?://\S+", "", l)
        if ":" in sin_url and not l.startswith(":::"):
            hallazgos.append(f"{nombre_md} línea {n} dos puntos, {l[:70]}")
        if re.search(r"[A-Za-zÁÉÍÓÚáéíóúñ]-[A-Za-zÁÉÍÓÚáéíóúñ]", sin_url):
            hallazgos.append(f"{nombre_md} línea {n} posible palabra compuesta con guion, {l[:70]}")
        if re.search(r"\b(delve|testament|tapestry|beacon|in conclusion)\b", l, re.I):
            hallazgos.append(f"{nombre_md} línea {n} cliché")
    return hallazgos


def main():
    datos = cargar_datos()
    nombre_archivo = re.sub(r"[^\w]+", "_", datos["NOMBRE"]).strip("_") or "software"
    if "PENDIENTE" in datos["NOMBRE"]:
        nombre_archivo = "PENDIENTE_NOMBRE"
    informe = []
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "referencia.docx"
        referencia_docx(ref)
        for fuente, salida, titulo in DOCUMENTOS:
            avisos = []
            texto = (FUENTES / fuente).read_text(encoding="utf-8")
            texto, numeros = resolver(texto, datos, avisos)
            md = Path(tmp) / fuente
            md.write_text(texto, encoding="utf-8")
            destino = DIR / salida.format(nombre=nombre_archivo)
            r = subprocess.run(["pandoc", str(md), "-f", "markdown+pipe_tables+raw_attribute+fenced_divs",
                                "-t", "docx", "--reference-doc", str(ref), "--resource-path", str(DIR),
                                "-o", str(destino)], capture_output=True, text=True, encoding="utf-8")
            if r.returncode != 0:
                print(r.stderr)
                sys.exit(1)
            posproceso(destino, datos["NOMBRE"], titulo)
            # Copia de las figuras con el número que tienen en el documento, para entregarlas aparte.
            carpeta_fig = DIR / "figuras_numeradas" / Path(fuente).stem
            carpeta_fig.mkdir(parents=True, exist_ok=True)
            for n, ruta in re.findall(r"!\[Figura (\d+)\.[^\]]*\]\(([^)]+)\)", texto):
                origen = DIR / ruta
                for ext in (".png", ".tif"):
                    fuente_ext = origen.with_suffix(ext)
                    if fuente_ext.exists():
                        shutil.copy2(fuente_ext, carpeta_fig / f"Figura_{int(n):02d}_{origen.stem}{ext}")
            informe.append(f"== {destino.name}")
            informe.append(f"figuras {len(numeros['fig'])}, tablas {len(numeros['tab'])}, "
                           f"ecuaciones {len(numeros['eq'])}, referencias {len(numeros['ref'])}")
            informe.extend(avisos)
            informe.extend(auditar(fuente, texto))
            if r.stderr.strip():
                informe.append("pandoc " + r.stderr.strip())
            print(f"{destino.name} escrito")
    (DIR / "auditoria_documentos.txt").write_text("\n".join(informe) + "\n", encoding="utf-8")
    print("\n".join(informe))


if __name__ == "__main__":
    main()
